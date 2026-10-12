from __future__ import annotations

import csv
import io
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

import httpx

from app.config import Settings
from app.connectors.base import Connector, SourceContext
from app.errors import UnsupportedSourceError
from app.logging import get_logger, log_event
from app.models.enums import SourcePlatform

CSV_SUFFIXES = (".csv", ".tsv", ".txt")
XLSX_SUFFIXES = (".xlsx", ".xlsm")
SNIFF_BYTES = 8192
MAX_SAMPLE_ROWS = 12

logger = get_logger(__name__)


class SpreadsheetConnector(Connector):
    platform = SourcePlatform.CSV

    def __init__(
        self,
        source: str,
        settings: Settings,
        client: httpx.AsyncClient | None = None,
    ) -> None:
        self.source = source
        self.settings = settings
        self.path = Path(source).expanduser()
        self._rows: list[dict[str, str]] | None = None
        self._headers: list[str] = []

    @property
    def pages_fetched(self) -> int:
        return 1 if self._rows is not None else 0

    async def discover(self) -> SourceContext:
        self._load()
        return SourceContext(
            platform=self._platform_for_suffix(),
            base_url=str(self.path),
            currency=None,
            store_name=self.path.stem,
            pages_fetched=self.pages_fetched,
        )

    async def iter_records(self) -> AsyncIterator[dict[str, Any]]:
        self._load()
        for row in self._rows or []:
            yield dict(row)

    @property
    def headers(self) -> list[str]:
        self._load()
        return list(self._headers)

    def sample(self, limit: int = MAX_SAMPLE_ROWS) -> list[dict[str, str]]:
        self._load()
        return [dict(row) for row in (self._rows or [])[:limit]]

    def _platform_for_suffix(self) -> SourcePlatform:
        return (
            SourcePlatform.XLSX
            if self.path.suffix.lower() in XLSX_SUFFIXES
            else SourcePlatform.CSV
        )

    def _load(self) -> None:
        if self._rows is not None:
            return
        if not self.path.is_file():
            raise UnsupportedSourceError(f"no such file: {self.path}")
        suffix = self.path.suffix.lower()
        if suffix in XLSX_SUFFIXES:
            self._rows, self._headers = _read_xlsx(self.path)
        elif suffix in CSV_SUFFIXES or not suffix:
            self._rows, self._headers = _read_delimited(self.path)
        else:
            raise UnsupportedSourceError(f"unsupported spreadsheet type: {suffix}")
        log_event(
            logger,
            "loaded spreadsheet",
            path=str(self.path),
            rows=len(self._rows),
            columns=len(self._headers),
        )


def _read_delimited(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if not text.strip():
        raise UnsupportedSourceError(f"{path} is empty")
    delimiter = _sniff(text)
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [h.strip() for h in reader.fieldnames or [] if h and h.strip()]
    if not headers:
        raise UnsupportedSourceError(f"{path} has no usable header row")
    rows: list[dict[str, str]] = []
    for raw in reader:
        row = {
            key.strip(): (value or "").strip()
            for key, value in raw.items()
            if key and key.strip()
        }
        if any(row.values()):
            rows.append(row)
    return rows, headers


def _sniff(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:SNIFF_BYTES], delimiters=",;\t|").delimiter
    except csv.Error:
        first = text.splitlines()[0] if text.splitlines() else ""
        counts = {d: first.count(d) for d in (",", ";", "\t", "|")}
        best = max(counts, key=lambda d: counts[d])
        return best if counts[best] else ","


def _read_xlsx(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UnsupportedSourceError(
            "reading .xlsx needs openpyxl, install it with: uv add openpyxl"
        ) from exc

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise UnsupportedSourceError(f"{path} has no readable sheet")
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise UnsupportedSourceError(f"{path} is empty") from exc
    headers = [str(value).strip() for value in header_row if value is not None]
    if not headers:
        raise UnsupportedSourceError(f"{path} has no usable header row")
    rows: list[dict[str, str]] = []
    for values in iterator:
        row = {
            header: ("" if value is None else str(value).strip())
            for header, value in zip(headers, values, strict=False)
        }
        if any(row.values()):
            rows.append(row)
    workbook.close()
    return rows, headers
